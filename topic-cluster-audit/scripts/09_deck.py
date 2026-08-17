#!/usr/bin/env python3
"""Generator 2 — the stakeholder deck.

Generated from the data, not written by hand. That's the whole point: a re-run refreshes
every number in every slide, so the deck can't drift from the workbook.

Output: out/deck.pptx
"""

import sys
from collections import defaultdict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from lib import OUT, load_config, read_csv

cfg = load_config()
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
DANGER = RGBColor(0xDC, 0x26, 0x26)

master = read_csv("master.csv")


def maybe(name):
    try:
        return read_csv(name)
    except SystemExit:
        return []


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    box = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(34), True, INK
    if kicker:
        k = tf.add_paragraph()
        k.text = kicker
        k.font.size, k.font.color.rgb = Pt(15), MUTED
    return s


def stat(s, left, value, label, color=INK):
    box = s.shapes.add_textbox(Inches(left), Inches(2.3), Inches(3.0), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(value)
    p.font.size, p.font.bold, p.font.color.rgb = Pt(48), True, color
    l = tf.add_paragraph()
    l.text = label
    l.font.size, l.font.color.rgb = Pt(13), MUTED


def bullets(s, items, top=2.4, size=16):
    box = s.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(4.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size, p.font.color.rgb = Pt(size), INK
        p.space_after = Pt(10)


def table(s, headers, rows, top=2.3, widths=None):
    if not rows:
        return
    shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.8), Inches(top),
                               Inches(11.7), Inches(0.4 * (len(rows) + 1)))
    t = shape.table
    if widths:
        total = sum(widths)
        for i, w in enumerate(widths):
            t.columns[i].width = Emu(int(Inches(11.7) * w / total))
    for c, h in enumerate(headers):
        cell = t.cell(0, c)
        cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        cell.text_frame.paragraphs[0].font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(v)
            cell.text_frame.paragraphs[0].font.size = Pt(10)


# --- Title -------------------------------------------------------------------
s = slide(f"{cfg['topic']} — topic cluster audit",
          f"{cfg['site_url']} · {len(master)} pages · generated from the pipeline, not by hand")

# --- The headline numbers ----------------------------------------------------
total_clicks = sum(int(r["clicks"] or 0) for r in master)
ranked = sorted(master, key=lambda r: -int(r["clicks"] or 0))
running, head = 0, 0
for r in ranked:
    running += int(r["clicks"] or 0)
    head += 1
    if total_clicks and running >= total_clicks * 0.93:
        break
retire = sum(1 for r in master if r["proposed_decision"] in ("KILL", "MERGE"))

s = slide("Where the cluster actually stands", "90 days of Search Console, joined to the page inventory")
stat(s, 0.9, len(master), "pages in the cluster")
stat(s, 4.0, f"{head / max(len(master), 1):.0%}", "of pages carry 93% of clicks", ACCENT)
stat(s, 7.1, retire, "pages proposed to retire or merge", DANGER)
stat(s, 10.2, f"{total_clicks:,}", "clicks, 90 days")

# --- Proposed decisions ------------------------------------------------------
tally = defaultdict(int)
for r in master:
    tally[r["proposed_decision"]] += 1
s = slide("Proposed calls", "The pipeline assembles evidence. Every call still gets reviewed by hand.")
table(s, ["Decision", "Pages", "What it means"],
      [[d, tally.get(d, 0), m] for d, m in [
          ("KEEP", "strong intent, no duplicate problem"),
          ("KEEP+REWRITE", "right intent, thin or overlapping"),
          ("MERGE", "a near-twin with a stronger page"),
          ("KILL", "the queries say nobody who matters is asking"),
          ("CHECK-FIRST", "may be in active campaigns — do not touch yet"),
      ]], widths=[2, 1, 6])

# --- The kill list -----------------------------------------------------------
kills = [r for r in master if r["proposed_decision"] == "KILL"][:12]
if kills:
    s = slide("Kill your winners when the queries say so",
              "Ranking top-10 for a term nobody who matters searches is not something to brag about.")
    table(s, ["Page", "Clicks", "Junk intent", "Why"],
          [[r["title"][:60], r["clicks"], f"{float(r.get('pct_junk') or 0):.0%}", r["why"][:50]] for r in kills],
          widths=[5, 1, 1.2, 4])

# --- Cannibalization ---------------------------------------------------------
dupes = maybe("similarity_pairs.csv")[:10]
if dupes:
    s = slide("Pages competing with each other",
              "Whole-page cosine catches twins. Max chunk cosine catches a duplicated middle section.")
    table(s, ["Page A", "Page B", "Page", "Max chunk", "Signal"],
          [[d["title_a"][:38], d["title_b"][:38], d["page_cosine"], d["max_chunk_cosine"], d["signal"]] for d in dupes],
          widths=[4, 4, 1, 1.2, 2])

# --- Expansion ---------------------------------------------------------------
exp = maybe("ahrefs_expansion.csv")[:10]
if exp:
    s = slide("Where competitors own a keyword and we have no page at all")
    table(s, ["Keyword", "Volume", "Competitor", "Their position"],
          [[e["keyword"][:55], e["volume"], e["competitor"], e["competitor_position"]] for e in exp],
          widths=[6, 1.4, 3, 1.6])

# --- Caveats -----------------------------------------------------------------
s = slide("Read this before you act on any of it")
bullets(s, [
    "The intent classification is only as good as the buckets. Read raw queries before you define them.",
    "Small embedding models are fine for near-duplicate detection. They are not fine for judging quality.",
    "The similarity threshold does not transfer between corpora. Recalibrate every time; it takes ten minutes.",
    "Search Console's per-URL query API caps what it returns and samples the long tail.",
    "Everything here is evidence for a human decision. Mistakes get through. Check the workbook.",
])

path = OUT / "deck.pptx"
prs.save(path)
print(f"wrote {path} ({len(prs.slides)} slides)", file=sys.stderr)
