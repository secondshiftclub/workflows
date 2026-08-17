#!/usr/bin/env python3
"""Generator 1 — the decision workbook.

A multi-sheet Excel file built from master.csv and its inputs. This is the artifact you
actually work in: one row per URL, sorted so the expensive decisions are at the top,
with the evidence sheets behind it.

Because it's generated from the data, a re-run refreshes every number.

Output: out/decisions.xlsx
"""

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from lib import OUT, load_config, read_csv

cfg = load_config()
DECISIONS = ["KEEP", "KEEP+REWRITE", "MERGE", "KILL", "CHECK-FIRST"]
COLORS = {
    "KEEP": "C6EFCE",
    "KEEP+REWRITE": "FFEB9C",
    "MERGE": "DDEBF7",
    "KILL": "FFC7CE",
    "CHECK-FIRST": "E4DFEC",
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def sheet(wb, name, rows, widths=None, freeze="A2"):
    ws = wb.create_sheet(name[:31])
    if not rows:
        ws["A1"] = "(empty — this step did not run)"
        return ws
    cols = list(rows[0].keys())
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_num(r.get(c, "")) for c in cols])
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, min(max(len(c) + 4, 12), 55))
    return ws


def _num(v):
    if isinstance(v, str) and v.strip():
        try:
            return int(v)
        except ValueError:
            try:
                return float(v)
            except ValueError:
                return v
    return v


master = read_csv("master.csv")


def maybe(name):
    try:
        return read_csv(name)
    except SystemExit:
        return []


wb = Workbook()
wb.remove(wb.active)

# --- Sheet 1: the summary a stakeholder reads -------------------------------
ws = wb.create_sheet("Summary")
ws["A1"] = f"{cfg['topic']} — topic cluster audit"
ws["A1"].font = Font(size=16, bold=True)
ws["A2"] = f"{len(master)} pages in the cluster · {cfg['site_url']}"
ws["A2"].font = Font(color="6B7280")

total_clicks = sum(int(r["clicks"] or 0) for r in master)
total_impr = sum(int(r["impressions"] or 0) for r in master)
ranked = sorted(master, key=lambda r: -int(r["clicks"] or 0))
running, head = 0, 0
for r in ranked:
    running += int(r["clicks"] or 0)
    head += 1
    if total_clicks and running >= total_clicks * 0.93:
        break

rows = [
    ("Pages in cluster", len(master)),
    ("Total clicks (90d)", total_clicks),
    ("Total impressions (90d)", total_impr),
    ("Pages carrying 93% of clicks", f"{head} ({head / max(len(master), 1):.0%})"),
    ("Thin pages", sum(1 for r in master if str(r.get("thin")).lower() == "true")),
    ("Near-duplicate pairs", len(maybe("similarity_pairs.csv"))),
    ("", ""),
]
for d in DECISIONS:
    rows.append((f"Proposed {d}", sum(1 for r in master if r["proposed_decision"] == d)))

for i, (k, v) in enumerate(rows, start=4):
    ws.cell(row=i, column=1, value=k).font = Font(bold=not k.startswith("Proposed"))
    ws.cell(row=i, column=2, value=v)
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 18

# --- Sheet 2: the one you work in -------------------------------------------
ws = sheet(wb, "Decisions", master, widths={"url": 55, "title": 40, "why": 45, "notes": 30})
if master:
    cols = list(master[0].keys())
    dcol = get_column_letter(cols.index("final_decision") + 1)
    pcol = get_column_letter(cols.index("proposed_decision") + 1)
    dv = DataValidation(type="list", formula1=f'"{",".join(DECISIONS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{dcol}2:{dcol}{len(master) + 1}")
    for i, r in enumerate(master, start=2):
        fill = COLORS.get(r["proposed_decision"])
        if fill:
            ws[f"{pcol}{i}"].fill = PatternFill("solid", fgColor=fill)

# --- Evidence sheets ---------------------------------------------------------
sheet(wb, "Duplicates", maybe("similarity_pairs.csv"), widths={"url_a": 50, "url_b": 50, "title_a": 35, "title_b": 35})
sheet(wb, "Thin pages", maybe("thin_pages.csv"), widths={"url": 55, "title": 40})
sheet(wb, "Page intent", maybe("page_intent.csv"), widths={"url": 55})
sheet(wb, "Query intent", maybe("query_intent.csv")[:5000], widths={"query": 50})
sheet(wb, "Expansion", maybe("ahrefs_expansion.csv"), widths={"keyword": 40, "competitor_url": 50})
sheet(wb, "Clusters", maybe("page_clusters.csv"), widths={"url": 55, "title": 40})

path = OUT / "decisions.xlsx"
wb.save(path)
print(f"wrote {path}", file=sys.stderr)
print(
    "Work in the Decisions sheet. final_decision is a dropdown; proposed_decision is\n"
    "colour-coded so you can see at a glance where the pipeline and you disagree.",
    file=sys.stderr,
)
